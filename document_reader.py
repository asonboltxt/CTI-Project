from pathlib import Path
from docx import Document
from pypdf import PdfReader

class DocumentReadError(Exception):
    pass

def clean_text(value):
    return " ".join((value or "").replace("\xa0", " ").split()).strip()

def read_docx(path):
    document = Document(str(path))
    paragraphs = [clean_text(p.text) for p in document.paragraphs if clean_text(p.text)]
    tables = []
    for table_index, table in enumerate(document.tables):
        rows = []
        for row_index, row in enumerate(table.rows):
            cells = [clean_text(cell.text) for cell in row.cells]
            if any(cells):
                rows.append({"row_index": row_index, "cells": cells})
        if rows:
            tables.append({"table_index": table_index, "rows": rows})
    text_rows = [" | ".join(row["cells"]) for table in tables for row in table["rows"]]
    return {
        "file_type": "docx",
        "text": "\n".join(paragraphs + text_rows),
        "paragraphs": paragraphs,
        "tables": tables,
        "warnings": [],
    }

def read_pdf(path):
    reader = PdfReader(str(path))
    pages = []
    warnings = []
    for index, page in enumerate(reader.pages):
        try:
            raw = page.extract_text(extraction_mode="layout") or ""
        except TypeError:
            raw = page.extract_text() or ""
        pages.append({"page": index + 1, "text": raw})
    text = "\n\f\n".join(page["text"] for page in pages)
    if len(clean_text(text)) < 100:
        warnings.append("The PDF has little extractable text and may require OCR.")
    return {
        "file_type": "pdf",
        "text": text,
        "paragraphs": [clean_text(x["text"]) for x in pages if clean_text(x["text"])],
        "tables": [],
        "pages": pages,
        "warnings": warnings,
    }

def read_document(path):
    path = Path(path)
    try:
        if path.suffix.lower() == ".docx":
            return read_docx(path)
        if path.suffix.lower() == ".pdf":
            return read_pdf(path)
    except Exception as exc:
        raise DocumentReadError(f"Unable to read {path.name}: {exc}") from exc
    raise DocumentReadError("Only DOCX and PDF files are supported.")
